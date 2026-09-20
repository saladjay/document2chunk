"""Excel 测试夹具：运行时用 openpyxl 生成 xlsx 字节（对齐 tests/_fixtures.py 的 PDF 夹具模式）。"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font


def build_xlsx(
    sheets: dict[str, list[list]],
    *,
    merges: tuple[str, ...] = (),          # 如 "A1:B1"
    bolds: dict[tuple[int, int], bool] = None,   # (row, col) 0-based
    number_formats: dict[tuple[int, int], str] = None,  # (row, col) 0-based
    hidden_cols: tuple[int, ...] = (),     # 0-based
    hidden_rows: tuple[int, ...] = (),     # 0-based
    hidden_sheets: tuple[str, ...] = (),
) -> bytes:
    bolds = bolds or {}
    number_formats = number_formats or {}
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                cell = ws.cell(row=r + 1, column=c + 1)
                if isinstance(v, str) and v.startswith("="):
                    cell.value = v  # 公式：openpyxl 不写缓存值 → 可测"无缓存"场景
                else:
                    cell.value = v
                if (r, c) in bolds:
                    cell.font = Font(bold=bolds[(r, c)])
                if (r, c) in number_formats:
                    cell.number_format = number_formats[(r, c)]
        for m in merges:
            ws.merge_cells(m)
        for c in hidden_cols:
            ws.column_dimensions[ws.cell(row=1, column=c + 1).column_letter].hidden = True
        for r in hidden_rows:
            ws.row_dimensions[r + 1].hidden = True
    for name in hidden_sheets:
        wb[name].sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
