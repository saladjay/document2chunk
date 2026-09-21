"""读取层：python-calamine 读值（快）+ openpyxl 二遍读格式（准）。

- calamine `skip_empty_area=True` 天然给出"实际扫描值域"（定档 #5：不信声明 dimension）。
- openpyxl 二遍只取：合并 ranges、数字格式、加粗、outline、隐藏行列、无缓存公式坐标、错误字面量坐标。
- 单元格数超 `_OPENPYXL_CELL_LIMIT` 时跳过二遍（性能护栏，定档 #18），并记 warning。
"""
from __future__ import annotations

import io
import zipfile

from document2chunk.extractors.excel.models import CellFmt, SheetGrid

_OPENPYXL_CELL_LIMIT = 200_000


def _zip_flags(data: bytes) -> dict[str, bool]:
    flags = {"external_links": False, "pivot": False}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            for n in zf.namelist():
                low = n.lower()
                if low.startswith("xl/externallinks/"):
                    flags["external_links"] = True
                elif low.startswith("xl/pivottables/") or low.startswith("xl/pivotcache/"):
                    flags["pivot"] = True
    except zipfile.BadZipFile:
        pass
    return flags


def _read_openpyxl_info(data: bytes, sheet_names: list[str]) -> dict[str, dict]:
    from openpyxl import load_workbook as _owlb

    wb_v = _owlb(io.BytesIO(data), data_only=True)   # 值/格式
    wb_f = _owlb(io.BytesIO(data), data_only=False)  # 公式坐标（data_only=True 时无缓存公式读不到 'f'）
    out: dict[str, dict] = {}
    for name in sheet_names:
        ws = wb_v[name]
        merged = [
            (rng.min_row - 1, rng.min_col - 1, rng.max_row - 1, rng.max_col - 1)
            for rng in ws.merged_cells.ranges
        ]
        formats: dict[tuple[int, int], CellFmt] = {}
        for row in ws.iter_rows():
            for cell in row:
                r, c = cell.row - 1, cell.column - 1
                if cell.data_type == "e":
                    formats[(r, c)] = CellFmt(
                        number_format=cell.number_format or "General",
                        bold=bool(cell.font and cell.font.bold),
                        outline_level=int(getattr(cell, "outline_level", 0) or 0),
                        error=str(cell.value),
                    )
                    continue
                if cell.value is None:
                    continue
                fmt = cell.number_format or "General"
                bold = bool(cell.font and cell.font.bold)
                outline = int(getattr(cell, "outline_level", 0) or 0)
                if fmt != "General" or bold or outline:
                    formats[(r, c)] = CellFmt(number_format=fmt, bold=bold, outline_level=outline)
        formula_no_cache: set[tuple[int, int]] = set()
        ws_f = wb_f[name]
        for row in ws_f.iter_rows():
            for cell in row:
                if cell.data_type == "f" and cell.row <= ws.max_row and cell.column <= ws.max_column:
                    v = ws.cell(row=cell.row, column=cell.column).value
                    if v in (None, ""):
                        formula_no_cache.add((cell.row - 1, cell.column - 1))
        hidden_rows = {i - 1 for i, dim in ws.row_dimensions.items() if dim.hidden}
        # column_dimensions 的键是列字母字符串（如 "A"），需用 column_index_from_string 转 1-based 再 -1
        from openpyxl.utils import column_index_from_string as _col_idx
        hidden_cols = {_col_idx(k) - 1 for k, dim in ws.column_dimensions.items() if dim.hidden}
        out[name] = {
            "state": ws.sheet_state,
            "merged": merged,
            "formats": formats,
            "hidden_rows": hidden_rows,
            "hidden_cols": hidden_cols,
            "formula_no_cache": formula_no_cache,
            "error_cells": {rc: f.error for rc, f in formats.items() if f.error},
        }
    return out


def read_sheet_grids(data: bytes) -> tuple[list[SheetGrid], list[str]]:
    """返回（可见 sheet 网格, warnings）。"""
    from python_calamine import load_workbook  # 延迟导入：excel extra

    warnings: list[str] = []
    flags = _zip_flags(data)
    if flags["external_links"]:
        warnings.append("文件含外部工作簿链接，相关单元格为保存时的缓存值")
    if flags["pivot"]:
        warnings.append("文件含数据透视表，其数据可能不在单元格中")

    wb = load_workbook(io.BytesIO(data))
    names = list(wb.sheet_names)
    sheet_rows = [wb.get_sheet_by_name(n).to_python(skip_empty_area=True) for n in names]
    total_cells = sum(len(r) for rows in sheet_rows for r in rows)

    info: dict[str, dict] = {}
    if total_cells <= _OPENPYXL_CELL_LIMIT:
        info = _read_openpyxl_info(data, names)
    else:
        warnings.append(
            f"单元格数 {total_cells} 超过上限 {_OPENPYXL_CELL_LIMIT}，跳过格式二遍（无合并/格式/隐藏信息）"
        )
    no_cache_total = sum(len(d.get("formula_no_cache", set())) for d in info.values())
    if no_cache_total:
        warnings.append(
            f"{no_cache_total} 个公式单元格无缓存值（文件可能由程序生成），读为空值"
        )

    grids: list[SheetGrid] = []
    for i, name in enumerate(names):
        d = info.get(name, {})
        if d.get("state", "visible") != "visible":
            warnings.append(f"隐藏 sheet「{name}」已跳过")
            continue
        grids.append(
            SheetGrid(
                name=name,
                values=[list(r) for r in sheet_rows[i]],
                merged=d.get("merged", []),
                formats=d.get("formats", {}),
                hidden_rows=d.get("hidden_rows", set()),
                hidden_cols=d.get("hidden_cols", set()),
                formula_no_cache=d.get("formula_no_cache", set()),
                error_cells=d.get("error_cells", {}),
                has_external_links=flags["external_links"],
                has_pivot=flags["pivot"],
            )
        )
    return grids, warnings
